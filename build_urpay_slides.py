#!/usr/bin/env python3
"""UrPay pitch deck — 2 English slides (16:9) for Zain Hackathon, Section 4.
Built natively with python-pptx: emerald/gold/paper palette, no blue/purple.
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ---------------- palette (brand: emerald / gold / paper / ink) ----------------
PAPER     = RGBColor(0xF8, 0xF6, 0xF0)   # warm paper off-white
EMERALD   = RGBColor(0x0E, 0x7A, 0x5C)   # primary
EMERALD_DK= RGBColor(0x0B, 0x5C, 0x46)   # dark emerald
GOLD      = RGBColor(0xCB, 0xA1, 0x35)   # warm gold (brand accent)
GOLD_LT   = RGBColor(0xE8, 0xC8, 0x67)   # light gold (on dark fills)
BRONZE    = RGBColor(0x8A, 0x6D, 0x1F)   # deep gold for small text on paper
GOLD_TINT = RGBColor(0xF1, 0xE9, 0xD4)   # gold-tinted surface
INK       = RGBColor(0x1B, 0x21, 0x1E)   # ink dark
MUTED     = RGBColor(0x4E, 0x58, 0x52)   # muted green-gray body text
HAIR      = RGBColor(0xDC, 0xD7, 0xC8)   # warm hairline
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
MINT      = RGBColor(0x7F, 0xC4, 0xA8)   # emerald-light (tags on dark)
CHAT_HAIR = RGBColor(0x3A, 0x42, 0x3D)
FONT      = "Arial"
SW, SH    = 13.333, 7.5

# ---------------- helpers ----------------
def rect(slide, x, y, w, h, fill=None, line=None, lw=1.0,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=None):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.shadow.inherit = False
    if adj is not None:
        try:
            sp.adjustments[0] = adj
        except Exception:
            pass
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(lw)
    return sp

def _runs(p, runs):
    for run in runs:
        txt, size, color, bold = run[0], run[1], run[2], run[3]
        extra = run[4] if len(run) > 4 else {}
        r = p.add_run(); r.text = txt
        f = r.font
        f.name = FONT; f.size = Pt(size); f.color.rgb = color; f.bold = bold
        if extra.get("spc"):
            f._rPr.set("spc", str(int(extra["spc"] * 100)))
        if extra.get("italic"):
            f.italic = True

def tbox(slide, x, y, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         line_sp=None, space_after=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line_sp: p.line_spacing = line_sp
        if space_after is not None: p.space_after = Pt(space_after)
        _runs(p, para)
    return tb

def stext(sp, paras, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE):
    tf = sp.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        _runs(p, para)
    return sp

def logo(slide, x, y, size, u_pt):
    """Rounded emerald square + gold 'U' + gold sun-dot."""
    sq = rect(slide, x, y, size, size, fill=EMERALD, adj=0.24)
    stext(sq, [[("U", u_pt, GOLD_LT, True)]])
    d = size * 0.13
    rect(slide, x + size * 0.74, y + size * 0.10, d, d, fill=GOLD_LT, shape=MSO_SHAPE.OVAL)
    return sq

def footer(slide, text, page):
    rect(slide, 0.62, 6.86, 12.09, 0.014, fill=HAIR, shape=MSO_SHAPE.RECTANGLE)
    tbox(slide, 0.62, 6.99, 9.5, 0.3, [[(text, 9.5, MUTED, False)]])
    tbox(slide, 11.21, 6.99, 1.5, 0.3, [[(page, 9.5, BRONZE, True)]], align=PP_ALIGN.RIGHT)

# ---------------- deck ----------------
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(SW), Inches(SH)
blank = prs.slide_layouts[6]

# ============ SLIDE 1 — "UrPay — Iraq's Agentic Payment Platform" ============
s1 = prs.slides.add_slide(blank)
rect(s1, 0, 0, SW, SH, fill=PAPER, shape=MSO_SHAPE.RECTANGLE)          # paper bg
rect(s1, 0, 0, 0.12, SH, fill=GOLD, shape=MSO_SHAPE.RECTANGLE)         # gold edge bar

# brand row: logo mark + wordmark + Arabic name
logo(s1, 0.62, 0.52, 0.66, 30)
tbox(s1, 1.46, 0.54, 3.4, 0.66,
     [[("Ur", 34, EMERALD, True), ("Pay", 34, GOLD, True)]],
     anchor=MSO_ANCHOR.MIDDLE)
tbox(s1, 11.05, 0.72, 1.66, 0.34, [[("أور پاي", 13, MUTED, False)]], align=PP_ALIGN.RIGHT)

# hero block (left column)
tbox(s1, 0.62, 1.72, 6.9, 0.3,
     [[("IRAQ'S AGENTIC PAYMENT PLATFORM", 12, BRONZE, True, {"spc": 1.6})]])
tbox(s1, 0.60, 2.06, 7.0, 1.7,
     [[("Every bill.", 48, INK, True)],
      [("One conversation.", 48, EMERALD, True)]], line_sp=1.02)
rect(s1, 0.62, 3.88, 1.3, 0.05, fill=GOLD, shape=MSO_SHAPE.RECTANGLE)  # gold rule
tbox(s1, 0.62, 4.10, 6.55, 0.7,
     [[("From Ur — where history's first transactions were recorded — to every Iraqi pocket.",
        12.5, MUTED, False)]], line_sp=1.15)

# problem card (right column)
rect(s1, 7.75, 1.70, 4.96, 1.38, fill=WHITE, line=HAIR, adj=0.07)
tbox(s1, 8.02, 1.92, 4.4, 0.26, [[("THE PROBLEM", 10, BRONZE, True, {"spc": 1.4})]])
tbox(s1, 8.02, 2.24, 4.42, 0.7,
     [[("Iraqis juggle ", 12, INK, False),
       ("6+ separate biller channels", 12, INK, True),
       (": queues, kiosks, and fragmented apps.", 12, INK, False)]], line_sp=1.12)

# solution card (emerald — the strong color block)
sc = rect(s1, 7.75, 3.26, 4.96, 1.62, fill=EMERALD, adj=0.07)
rect(s1, 8.02, 3.52, 0.10, 0.10, fill=GOLD_LT, shape=MSO_SHAPE.OVAL)
tbox(s1, 8.22, 3.47, 4.2, 0.26, [[("THE SOLUTION", 10, WHITE, True, {"spc": 1.4})]])
tbox(s1, 8.02, 3.82, 4.44, 0.95,
     [[("One wallet", 12, WHITE, True),
       (" + an ", 12, PAPER, False),
       ("AI agent", 12, WHITE, True),
       (" that pays electricity, water, internet, mobile, education & traffic bills — from a "
        "single chat, secured by PIN.", 12, PAPER, False)]], line_sp=1.12)

# 4 mini stat chips
chips = ["6 bill categories", "PIN-secured payments", "Instant P2P transfers", "Groq gpt-oss-120b"]
cx = 0.62
for c in chips:
    ch = rect(s1, cx, 5.66, 2.91, 0.56, fill=WHITE, line=GOLD, lw=1.25, adj=0.5)
    stext(ch, [[(c, 11, INK, True)]])
    cx += 2.91 + 0.15

footer(s1, "Zain Hackathon · Section 4 — Bill Pay Agent · 2026", "01")

# ============ SLIDE 2 — "How Ur Works — The Agentic Payment Loop" ============
s2 = prs.slides.add_slide(blank)
rect(s2, 0, 0, SW, SH, fill=PAPER, shape=MSO_SHAPE.RECTANGLE)
rect(s2, 0, 0, 0.12, SH, fill=GOLD, shape=MSO_SHAPE.RECTANGLE)

# header
tbox(s2, 0.62, 0.42, 6.0, 0.26, [[("HOW UR WORKS", 11, BRONZE, True, {"spc": 1.6})]])
tbox(s2, 0.62, 0.70, 8.2, 0.6, [[("The Agentic Payment Loop", 30, INK, True)]])
rect(s2, 0.62, 1.36, 0.95, 0.05, fill=GOLD, shape=MSO_SHAPE.RECTANGLE)
logo(s2, 11.02, 0.40, 0.36, 16)
tbox(s2, 11.48, 0.40, 1.3, 0.36,
     [[("Ur", 15, EMERALD, True), ("Pay", 15, GOLD, True)]], anchor=MSO_ANCHOR.MIDDLE)

# 4-step agent loop
steps = [
    ("01", "Understand", [[("Parses Arabic / English intent.", 10, MUTED, False)]]),
    ("02", "Verify", [[("Pulls real bills & balance via tools.", 10, MUTED, False)]]),
    ("03", "Authorize", [[("Requests the user's PIN — ", 10, MUTED, False),
                          ("never stored", 10, EMERALD_DK, True)]]),
    ("04", "Execute & Receipt", [[("Payment + ", 10, MUTED, False),
                                  ("UR-reference receipt", 10, EMERALD_DK, True)]]),
]
CW, CH, GAP = 2.72, 1.98, 0.4033
for i, (num, title, desc) in enumerate(steps):
    x = 0.62 + i * (CW + GAP)
    rect(s2, x, 1.62, CW, CH, fill=WHITE, line=HAIR, adj=0.075)
    chip = rect(s2, x + 0.22, 1.86, 0.40, 0.40, fill=EMERALD, adj=0.30)
    stext(chip, [[(num, 12.5, WHITE, True)]])
    tbox(s2, x + 0.22, 2.40, CW - 0.40, 0.34, [[(title, 15.5, INK, True)]])
    tbox(s2, x + 0.22, 2.76, CW - 0.42, 0.75, desc, line_sp=1.12)
    if i < 3:  # gold chevron arrow between steps
        ax = x + CW + (GAP - 0.34) / 2
        rect(s2, ax, 2.40, 0.34, 0.42, fill=GOLD, shape=MSO_SHAPE.CHEVRON)

# tech stack chips (dark pills)
tbox(s2, 0.62, 3.96, 4.0, 0.24, [[("TECH STACK", 10, BRONZE, True, {"spc": 1.4})]])
tech = [("FastAPI", 0.98), ("SQLite", 0.90), ("Next.js 16", 1.16),
        ("Groq gpt-oss-120b", 1.74), ("Function Calling", 1.62)]
tx = 0.62
for label, w in tech:
    c = rect(s2, tx, 4.24, w, 0.42, fill=INK, adj=0.5)
    stext(c, [[(label, 9.5, PAPER, True)]])
    tx += w + 0.14

# security & data note (gold-tint box)
rect(s2, 0.62, 4.96, 7.0, 1.14, fill=GOLD_TINT, adj=0.10)
tbox(s2, 0.88, 5.12, 4.0, 0.24, [[("SECURITY & DATA", 9.5, BRONZE, True, {"spc": 1.4})]])
tbox(s2, 0.88, 5.40, 6.48, 0.6,
     [[("PBKDF2-hashed PIN", 11, INK, True),
       (" · JWT sessions · masked PIN in chat logs · ", 11, INK, False),
       ("100 realistic seeded Iraqi users", 11, INK, True)]], line_sp=1.12)

# chat mockup (dark rounded card, 3 lines)
rect(s2, 7.9, 3.92, 4.81, 2.55, fill=INK, adj=0.075)
rect(s2, 8.16, 4.17, 0.10, 0.10, fill=GOLD_LT, shape=MSO_SHAPE.OVAL)
tbox(s2, 8.36, 4.12, 3.9, 0.24, [[("Ur — Agent Chat", 9.5, GOLD_LT, True, {"spc": 1.0})]])
rect(s2, 8.16, 4.48, 4.28, 0.014, fill=CHAT_HAIR, shape=MSO_SHAPE.RECTANGLE)
tbox(s2, 8.16, 4.66, 4.30, 1.66, [
    [("USER", 8, GOLD_LT, True, {"spc": 1.2}), ("   ", 8, GOLD_LT, True),
     ("ادفع فاتورة الكهرباء", 11, PAPER, False)],
    [("UR", 8, MINT, True, {"spc": 1.2}), ("   ", 8, MINT, True),
     ("Electricity bill, ", 9.5, PAPER, False), ("45,000 IQD", 9.5, GOLD_LT, True),
     (". Send your PIN.", 9.5, PAPER, False)],
    [("UR", 8, MINT, True, {"spc": 1.2}), ("   ", 8, MINT, True),
     ("✅ Paid — receipt ", 9.5, PAPER, False), ("UR-8XK2F3", 9.5, GOLD_LT, True),
     (", balance ", 9.5, PAPER, False), ("1,705,000 IQD", 9.5, PAPER, False)],
], space_after=7)

footer(s2, "Demo card 4539 1234 1234 1234 · PIN 1234 — urpay.iq (demo)", "02")

# ---------------- save ----------------
prs.core_properties.title = "UrPay — Iraq's Agentic Payment Platform"
prs.core_properties.author = "UrPay Team"
prs.core_properties.subject = "Zain Hackathon · Section 4 — Bill Pay Agent"
prs.core_properties.comments = "Every bill. One conversation."

OUT = "/home/z/my-project/public/UrPay-Slides.pptx"
os.makedirs(os.path.dirname(OUT), exist_ok=True)
prs.save(OUT)
print("saved:", OUT, os.path.getsize(OUT), "bytes")
